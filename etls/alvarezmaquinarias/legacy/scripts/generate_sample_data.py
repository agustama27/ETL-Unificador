"""Genera los 4 archivos fuente fake en inputs/<run-date>/ para probar el
ETL end-to-end sin los archivos reales de Álvarez Maquinaria.

Los datos son ficticios pero replican la *estructura sucia* real de los
exports de Autologica (header corrido, ID de cliente entre paréntesis,
filas de subtotal de pivot, PDF de repuestos) mapeada en
docs/ASSUMPTIONS.md y docs/SPEC.md — no la vieja estructura plana.

Uso: python scripts/generate_sample_data.py [--run-date AAAA-MM-DD]
"""

from __future__ import annotations

import argparse
import csv
from datetime import date
from pathlib import Path
from typing import Sequence

import xlwt
from openpyxl import Workbook
from reportlab.pdfgen import canvas

PROJECT_ROOT = Path(__file__).resolve().parent.parent


def _write_xls(path: Path, rows: list[list]) -> None:
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet1")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if value is None:
                continue
            ws.write(r, c, value)
    wb.save(str(path))


def _write_xlsx(path: Path, sheets: dict[str, list[list]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(str(path))


def _saldos_generales_rows() -> list[list]:
    """Header corrido (fila 3, no fila 0) + filas de título antes, ID
    Autologica entre paréntesis junto al nombre. Ver docs/ASSUMPTIONS.md
    sección 1.1."""
    return [
        ["Saldos de clientes"],
        ["Filtrado por:", "Fecha hasta 2/07/2026; Segun saldo > 30"],
        [],
        ["Cliente", None, "CUIT/CUIL", None, "Localidad", "Telefono", None, "Saldo DOLARES", "Ultimo comprobante", "Fecha"],
        ["AGRO GANADERA FAKE S.A.", "(101)", "30-1-1", None, "Localidad X", "3537411234", None, 150000.00, "CDNR1", "15/07/2026"],
        ["AGROPECUARIA DEL SUR", "(102)", "30-2-2", None, "Localidad Y", "3516541234", None, 25000.00, "CDNR2", "20/06/2026"],
        ["LOGISTICA BELL VILLE SRL", "(103)", "30-3-3", None, "Localidad Z", "0", None, 10000.00, "CDNR3", "10/07/2026"],
        ["CLIENTE SIN TELEFONO SA", "(104)", "30-4-4", None, "Localidad W", None, None, 50000.00, "CDNR4", "01/07/2026"],
        # mismo telefono que AGROPECUARIA DEL SUR pero razon social distinta:
        # normalize_client_name NO los une (no es variante del mismo nombre),
        # asi que son dos filas del Roman compartiendo linea -> el E1KIA debe
        # listar ese numero una sola vez
        ["CAMPO VERDE SRL", "(106)", "30-6-6", None, "Localidad Y", "3516541234", None, 7000.00, "CDNR7", "18/06/2026"],
        # nombre completo (40 car) del que servicios trae cortado a 27
        ["CONSTRUCTORA LOS ALAMOS SOCIEDAD ANONIMA", "(107)", "30-7-7", None, "Localidad Q", "3512001122", None, 4000.00, "CDNR8", "12/07/2026"],
        # par prefijo pero NINGUNO truncado: persona y sociedad de hecho son
        # deudores distintos y no deben fusionarse
        ["PEREZ JUAN", "(108)", "30-8-8", None, "Localidad R", "3513004455", None, 3000.00, "CDNR9", "05/07/2026"],
        ["PEREZ JUAN Y HERMANOS SH", "(109)", "30-9-9", None, "Localidad R", "3513009988", None, 2000.00, "CDNR10", "06/07/2026"],
        # misma persona que "GOMES ANDRES" de maquinarias: errata de una letra
        # y mismo telefono -> se fusionan
        ["GOMEZ ANDRES", "(110)", "30-10-10", None, "Localidad S", "3514007766", None, 900.00, "CDNR11", "08/07/2026"],
        # fila sin cliente (cierre de reporte): se ignora, no cuenta como descarte
        [None, None, None, None, None, "3537000000", None, 1000.00, "CDNR5", "01/01/2026"],
        # fila invalida a proposito: saldo no parseable -> debe descartarse y contarse
        ["CLIENTE MONTO INVALIDO SA", "(105)", "30-5-5", None, "Localidad V", "3512223344", None, "no-es-numero", "CDNR6", "01/01/2026"],
    ]


def _saldos_csv_sections() -> list[list[list[str]]]:
    """Reporte CSV sintético con dos secciones completas y repetidas."""
    header = [
        "Cliente",
        "ID Autologica",
        "Telefono",
        "Saldo DOLARES",
        "Ultimo comprobante",
        "Fecha",
    ]
    return [
        [
            ["Saldos de clientes"],
            header,
            ["AGRO GANADERA FAKE S.A.", "(101)", "3537411234", "150000,00", "CDNR1", "15/07/2026"],
        ],
        [
            ["Saldos de clientes"],
            header,
            ["AGROPECUARIA DEL SUR", "(102)", "3516541234", "25000,00", "CDNR2", "20/06/2026"],
        ],
    ]


def _write_saldos_csv(path: Path, sections: list[list[list[str]]] | None = None) -> None:
    """Escribe únicamente el CSV CP1252/coma soportado por el adapter."""
    with path.open("w", encoding="cp1252", newline="") as csv_file:
        writer = csv.writer(csv_file, delimiter=",", lineterminator="\r\n")
        for section in sections or _saldos_csv_sections():
            writer.writerows(section)


def _maquinarias_rows() -> list[list]:
    """Header corrido (fila 1, no fila 0), teléfono alternativo como
    fallback. Ver docs/ASSUMPTIONS.md sección 1.2."""
    return [
        [],
        ["CLIENTES", "UNIDAD", "DEUDA", "ULTIMO PAGO", "USD", "TELEFONO", None, "TEL. ALTERNATIVO", None],
        ["AGRO GANADERA FAKE S.A.", "NH 8030", 2017, 2021, 12000.00, None, None, None, None],
        ["ESTABLECIMIENTO EL EJEMPLO", "Tractor Fake 200", 2019, 2023, 8000.00, "invalido!!", None, "3537555666", None],
        # errata de la fuente: es el "GOMEZ ANDRES" de saldos, mismo telefono
        ["GOMES ANDRES", "Sembradora Fake", 2018, 2022, 5000.00, "3514007766", None, None, None],
        # fila invalida a proposito: USD no parseable -> debe descartarse y contarse
        ["CLIENTE MONTO INVALIDO MAQ", "Equipo X", 2020, 2020, "no-es-numero", None, None, None, None],
    ]


def _remitos_servicios_sheets() -> dict[str, list[list]]:
    """Dos hojas con layout de columnas distinto entre sí (ver
    docs/SPEC.md sección 2.3): `Table 1` con remito+cliente pegados en
    una celda (truncado) y precedidos por su fila de subtotal; `Table 2`
    con remito y cliente en columnas separadas, sin subtotal."""
    table1 = [
        ["N Remito                                         Cliente", "Fecha", "Precio", "IVA 21%", "Total"],
        ["Cliente: AGROPECUARIA DEL SUR\nRecuento: 1", "Max.: 20/06/2026", "Suma: USD320,00", "Suma: USD67,20", "Suma: USD387,20"],
        ["9011                                               AGROPECUARIA DEL SUR", date(2026, 6, 20), 320.00, 67.20, 387.20],
        ["Cliente: MANTENIMIENTOS CORDOBA\nRecuento: 1", "Max.: 15/07/2026", "Suma: USD1.500,00", "Suma: USD315,00", "Suma: USD1.815,00"],
        ["9012                                               MANTENIMIENTOS CORDOBA", date(2026, 7, 15), 1500.00, 315.00, 1815.00],
    ]
    table2 = [
        [None, "N Remito", "Cliente", "Fecha", "Precio", "IVA 21%", "Total"],
        [9013, "LOGISTICA BELL VILLE SRL", None, date(2026, 4, 5), 45.00, 9.45, 54.45],
        [9010, "SERVICIOS INDUSTRIALES S.H.", None, date(2026, 5, 12), 850.00, 178.50, 1028.50],
        # segundo remito de un cliente que ya aparece en Table 1: la fuente no
        # distingue los conceptos, ambos entran como "Servicios" y deben
        # consolidarse en una sola entrada del array Productos
        [9014, "AGROPECUARIA DEL SUR", None, date(2026, 7, 10), 200.00, 42.00, 242.00],
        # cortado por el ancho de columna del pivot (27 car, el techo de esta
        # fuente): es "CONSTRUCTORA LOS ALAMOS SOCIEDAD ANONIMA" de saldos
        [9015, "CONSTRUCTORA LOS ALAMOS SOC", None, date(2026, 6, 30), 500.00, 105.00, 605.00],
    ]
    return {"Table 1": table1, "Table 2": table2}


def _write_remitos_repuestos_pdf(path: Path, lines: list[str] | None = None) -> None:
    """Texto plano por línea replicando el patrón real (ver
    docs/SPEC.md sección 2.3): remito+cliente pegados a `REPUESTOS`, un
    cliente con factura ARS (sin flag dólar), que el extractor debe excluir,
    y una línea corrupta para testear el descarte seguro."""
    lines = lines or [
        "Remito Cliente Detalles Fecha Precio IVA 21% Total Factura relacionadaColumna 2",
        "Cliente: AGRO GANADERA FAKE S.A.Recuento: 1",
        "1501 AGRO GANADERA FAKE S.A.REPUESTOS 15/07/2026  $450,50  $94,60  $545,10PRECIO EN DOLARES",
        "Cliente: DISTRIBUIDORA NORTE SRLRecuento: 1",
        "1503 DISTRIBUIDORA NORTE SRLREPUESTOS 05/05/2026  $85,00  $17,85  $102,85Factura relacionada",
        "Cliente: AGRICOLA LAS LUNAS SASRecuento: 1",
        # linea corrupta a proposito: monto sin separador decimal -> no matchea el regex, debe descartarse
        "1504 AGRICOLA LAS LUNAS SASREPUESTOS 22/07/2026  $3500$735,00  $4235,00PRECIO EN DOLARES",
    ]
    c = canvas.Canvas(str(path), pagesize=(842, 595))
    text = c.beginText(20, 570)
    text.setFont("Helvetica", 9)
    for line in lines:
        text.textLine(line)
    c.drawText(text)
    c.showPage()
    c.save()


def generate_sample_data(project_root: Path, run_date: date) -> Path:
    """Escribe fixtures fake con los nombres fijos de una corrida fechada."""
    input_dir = project_root / "inputs" / run_date.isoformat()
    input_dir.mkdir(parents=True, exist_ok=True)
    _write_xls(input_dir / "saldos.xls", _saldos_generales_rows())
    _write_xlsx(input_dir / "maquinarias.xlsx", {"Sheet1": _maquinarias_rows()})
    _write_xlsx(input_dir / "servicios.xlsx", _remitos_servicios_sheets())
    _write_remitos_repuestos_pdf(input_dir / "repuestos.pdf")
    return input_dir


def _parse_run_date(value: str) -> date:
    try:
        parsed = date.fromisoformat(value)
    except ValueError as error:
        raise argparse.ArgumentTypeError("usar fecha AAAA-MM-DD válida") from error
    if parsed.isoformat() != value:
        raise argparse.ArgumentTypeError("usar fecha canónica AAAA-MM-DD")
    return parsed


def main(argv: Sequence[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-date", type=_parse_run_date, default=date.today())
    args = parser.parse_args(argv)
    input_dir = generate_sample_data(PROJECT_ROOT, args.run_date)
    print(f"Datos fake generados en {input_dir}")


if __name__ == "__main__":
    main()
